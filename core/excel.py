"""Excel export helpers (openpyxl - MIT, https://foss.heptapod.net/openpyxl/openpyxl).

Every sheet produced here contains quantities only. There is deliberately no
price, rate, amount, tax or value column anywhere in this module.
"""
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from django.http import HttpResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEAD_FILL = PatternFill("solid", fgColor="10243D")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
TITLE_FONT = Font(bold=True, size=13)
META_FONT = Font(size=9, color="666666")
THIN = Side(style="thin", color="D8DEE7")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
TOTAL_FONT = Font(bold=True)
TOTAL_FILL = PatternFill("solid", fgColor="EEF1F5")


def _cast(value):
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (datetime, date)):
        return value
    if value is None:
        return ""
    return value


def build_workbook(sheets):
    """sheets = [{'name','title','subtitle','columns','rows','totals'}] -> Workbook.

    columns = [(header, key_or_index, 'num'|'text'|'date'), ...] or plain headers.
    """
    wb = Workbook()
    wb.remove(wb.active)
    for spec in sheets:
        ws = wb.create_sheet(spec["name"][:31])
        columns = spec["columns"]
        headers = [c[0] if isinstance(c, (list, tuple)) else c for c in columns]
        kinds = [(c[2] if isinstance(c, (list, tuple)) and len(c) > 2 else "text")
                 for c in columns]
        ncols = len(headers)

        row_i = 1
        if spec.get("title"):
            ws.cell(1, 1, spec["title"]).font = TITLE_FONT
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(ncols, 1))
            row_i = 2
        if spec.get("subtitle"):
            ws.cell(row_i, 1, spec["subtitle"]).font = META_FONT
            ws.merge_cells(start_row=row_i, start_column=1, end_row=row_i,
                           end_column=max(ncols, 1))
            row_i += 1
        row_i += 1

        header_row = row_i
        for c, head in enumerate(headers, start=1):
            cell = ws.cell(header_row, c, head)
            cell.fill, cell.font, cell.border = HEAD_FILL, HEAD_FONT, BORDER
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[header_row].height = 26

        widths = [max(10, min(42, len(str(h)) + 3)) for h in headers]
        r = header_row
        for data_row in spec["rows"]:
            r += 1
            for c, value in enumerate(data_row[:ncols], start=1):
                cell = ws.cell(r, c, _cast(value))
                cell.border = BORDER
                kind = kinds[c - 1]
                if kind == "num":
                    cell.alignment = Alignment(horizontal="right")
                    cell.number_format = "#,##0.###"
                elif kind == "date":
                    cell.number_format = "DD-MM-YYYY"
                    cell.alignment = Alignment(horizontal="center")
                widths[c - 1] = max(widths[c - 1], min(48, len(str(value)) + 2))

        if spec.get("totals"):
            r += 1
            for c, value in enumerate(spec["totals"][:ncols], start=1):
                cell = ws.cell(r, c, _cast(value))
                cell.font, cell.fill, cell.border = TOTAL_FONT, TOTAL_FILL, BORDER
                if kinds[c - 1] == "num":
                    cell.alignment = Alignment(horizontal="right")
                    cell.number_format = "#,##0.###"

        for c, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(c)].width = w
        ws.freeze_panes = ws.cell(header_row + 1, 1)
        ws.auto_filter.ref = f"A{header_row}:{get_column_letter(max(ncols, 1))}{r}"
    return wb


def workbook_response(wb, filename):
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    response = HttpResponse(
        buf.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    return response


def excel_response(sheets, filename):
    return workbook_response(build_workbook(sheets), filename)
